"""Standalone isolated parser process. Never import this into the core runtime."""
import hashlib
import importlib.metadata
import io
import json
from pathlib import PurePosixPath, Path
import sys
import zipfile

PINS = {"pypdf": "6.20.1", "python-docx": "1.2.0", "openpyxl": "3.1.5",
        "lxml": "6.1.3", "et-xmlfile": "2.0.0", "typing_extensions": "4.15.0"}
TEXT_LIMIT = 524288
LOCATOR_LIMIT = 10000


class Rejected(Exception):
    pass


def archive_guard(content):
    if content.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        # OOXML encryption uses an OLE envelope; legacy Office also uses OLE.
        raise Rejected("encrypted_or_legacy_office")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 4096:
            raise Rejected("archive_limit")
        total = 0
        names = set()
        for entry in entries:
            name = entry.filename
            if ("\\" in name or ":" in name or "\x00" in name or
                    PurePosixPath(name).is_absolute() or ".." in PurePosixPath(name).parts or
                    name in names or (entry.external_attr >> 16) & 0o170000 == 0o120000):
                raise Rejected("unsafe_archive_path")
            names.add(name)
            if entry.flag_bits & 1:
                raise Rejected("encrypted_archive")
            total += entry.file_size
            if (total > 67108864 or entry.file_size > 16777216 or
                    entry.file_size > max(1, entry.compress_size) * 200):
                raise Rejected("archive_limit")
            if entry.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                raise Rejected("unsupported_archive_compression")
        # Package relationships can point to XML with any filename or extension.
        # Conservatively inspect every bounded entry before third-party parsing.
        for entry in entries:
            data = archive.read(entry)
            normalized = data.replace(b"\x00", b"").upper()
            if b"<!DOCTYPE" in normalized or b"<!ENTITY" in normalized:
                raise Rejected("unsafe_xml")
        return names


def extract(content, mime):
    text = ""
    locators = []
    partial = False
    empty = 0
    extractor = "optional-1"
    coverage = {"scope": "unassessed"}

    def finish(status, reason):
        return dict(status=status, reason=reason, text=text, locators=locators,
                    output_hash=hashlib.sha256(text.encode()).hexdigest(), extractor=extractor,
                    coverage=coverage)

    def append(value, locator):
        nonlocal text, partial
        if not value:
            return True
        separator = "\n" if text else ""
        remaining = TEXT_LIMIT - len(text) - len(separator)
        if remaining <= 0 or len(locators) >= LOCATOR_LIMIT:
            partial = True
            return False
        start = len(text) + len(separator)
        text += separator + value[:remaining]
        locators.append(dict(locator, start=start, end=len(text)))
        if len(value) > remaining:
            partial = True
            return False
        return True

    try:
        required = ({"pypdf"} if mime == "application/pdf" else
                    {"python-docx", "lxml", "typing_extensions"} if "wordprocessingml" in mime else
                    {"openpyxl", "et-xmlfile"})
        try:
            if any(importlib.metadata.version(name) != PINS[name] for name in required):
                return finish("unsupported", "dependency_version")
        except importlib.metadata.PackageNotFoundError:
            return finish("unsupported", "missing_dependency")
        if mime == "application/pdf":
            from pypdf import PdfReader
            extractor = "pypdf/" + PINS["pypdf"]
            coverage = dict(scope="text_layer", total_pages=None, visited_pages=0,
                            pages_with_text=0, missing_text_pages=[])
            document = PdfReader(io.BytesIO(content), strict=True)
            if document.is_encrypted:
                return finish("encrypted", "encrypted_pdf")
            coverage["total_pages"] = len(document.pages)
            for number, page in enumerate(document.pages, 1):
                if number > LOCATOR_LIMIT:
                    partial = True
                    break
                value = page.extract_text() or ""
                coverage["visited_pages"] += 1
                if not value.strip():
                    empty += 1
                    coverage["missing_text_pages"].append(number)
                else:
                    coverage["pages_with_text"] += 1
                if not append(value, {"page": number}):
                    break
            if empty or not text:
                return finish("partial", "incomplete_text_layer" if text else "no_text_layer")
        else:
            names = archive_guard(content)
            if "wordprocessingml" in mime:
                from docx import Document
                extractor = "python-docx/" + PINS["python-docx"]
                document = Document(io.BytesIO(content))
                paragraphs = document.paragraphs
                omissions = {"tables": len(document.tables)}
                for key, expression in {
                        "drawings": ".//w:drawing | .//w:pict", "embedded_objects": ".//w:object",
                        "tracked_changes": ".//w:ins | .//w:del", "notes": ".//w:footnoteReference | .//w:endnoteReference",
                        "alt_chunks": ".//w:altChunk", "content_controls": ".//w:sdt"}.items():
                    omissions[key] = len(document.part.element.xpath(expression))
                for key, prefix in (("headers", "word/header"), ("footers", "word/footer")):
                    omissions[key] = sum(name.startswith(prefix) and name.endswith(".xml") for name in names)
                coverage = dict(scope="body_paragraphs", total_paragraphs=len(paragraphs),
                                visited_paragraphs=0, nonempty_paragraphs=0, omissions=omissions)
                for number, paragraph in enumerate(paragraphs, 1):
                    coverage["visited_paragraphs"] += 1
                    coverage["nonempty_paragraphs"] += bool(paragraph.text)
                    if not append(paragraph.text, {"paragraph": number}):
                        break
                # Tables, headers, drawings and notes are outside this adapter's coverage.
                partial = partial or any(omissions.values())
            else:
                from openpyxl import load_workbook
                extractor = "openpyxl/" + PINS["openpyxl"]
                document = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
                omissions = {key: sum(name.startswith(prefix) and not name.endswith("/") for name in names)
                             for key, prefix in (("charts", "xl/charts/"), ("images", "xl/media/"),
                                                 ("external_links", "xl/externalLinks/"), ("comments", "xl/comments"))}
                omissions["macros"] = sum(name.endswith("vbaProject.bin") for name in names)
                coverage = dict(scope="stored_cells", total_sheets=len(document.sheetnames),
                                visited_sheets=0, visited_cells=0, nonempty_cells=0, omissions=omissions)
                try:
                    visited = 0
                    for sheet in document:
                        coverage["visited_sheets"] += 1
                        sheet.reset_dimensions()
                        for row in sheet.iter_rows():
                            for cell in row:
                                visited += 1
                                if visited > 100000:
                                    partial = True
                                    break
                                if not hasattr(cell, "coordinate"):
                                    continue
                                coverage["visited_cells"] += 1
                                coverage["nonempty_cells"] += cell.value is not None
                                if not append(str(cell.value) if cell.value is not None else "",
                                                                 {"sheet": sheet.title, "cell": cell.coordinate} if cell.value is not None else {}):
                                    partial = True
                                    break
                            if partial:
                                break
                        if partial:
                            break
                finally:
                    document.close()
                partial = partial or any(omissions.values())
        return finish("partial" if partial else "complete", "coverage_limit" if partial else "")
    except ImportError:
        text, locators = "", []
        return finish("unsupported", "missing_dependency")
    except Rejected as error:
        text, locators = "", []
        return finish("encrypted" if str(error) == "encrypted_archive" else "failed", str(error))
    except Exception:
        text, locators = "", []
        return finish("failed", "malformed_document")


if __name__ == "__main__":
    try:
        # Best effort POSIX limits supplement the parent's portable wall/output bounds.
        try:
            import resource
            resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
            resource.setrlimit(resource.RLIMIT_AS, (1073741824, 1073741824))
        except (ImportError, OSError, ValueError):
            pass
        with Path(sys.argv[1]).open("rb") as source:
            content = source.read(33554433)
        if len(content) > 33554432:
            raise ValueError("size_limit")
        value = extract(content, sys.argv[2])
        sys.stdout.write(json.dumps(value, ensure_ascii=True))
    except Exception:
        sys.stdout.write(json.dumps(dict(status="failed", reason="worker_failed", text="", locators=[],
                                       output_hash=hashlib.sha256(b"").hexdigest(), extractor="optional-1",
                                       coverage={"scope": "unassessed"})))
