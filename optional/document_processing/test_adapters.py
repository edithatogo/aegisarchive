"""Run under the explicitly provisioned venv for real parser acceptance."""
import hashlib
import io
import os
from pathlib import Path
import struct
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from . import runner
from . import worker

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def archive(entries, compression=zipfile.ZIP_STORED):
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w", compression=compression) as output:
        for name, value in entries:
            output.writestr(name, value)
    return target.getvalue()


class ParentTests(unittest.TestCase):
    def test_parser_suite_requires_exact_pins(self):
        with patch.object(worker.importlib.metadata, "version", return_value="0.0.0"):
            with self.assertRaises(unittest.SkipTest):
                ParserTests.setUpClass()

    def test_dependency_rejection(self):
        for mime in (PDF, DOCX, XLSX):
            with patch.object(worker.importlib.metadata, "version", return_value="0.0.0"):
                value = worker.extract(b"synthetic", mime)
                self.assertEqual((value["status"], value["reason"]), ("unsupported", "dependency_version"))
                self.assertEqual(value["coverage"], {"scope": "unassessed"})
            with patch.object(worker.importlib.metadata, "version", side_effect=worker.importlib.metadata.PackageNotFoundError):
                self.assertEqual(worker.extract(b"synthetic", mime)["reason"], "missing_dependency")

    def test_malformed_worker_schemas(self):
        payloads = [None, [], "invalid", {}, dict(runner.result("complete", ""), status=[]),
                    dict(runner.result("complete", ""), coverage=[]),
                    dict(runner.result("complete", ""), coverage={"scope": []}),
                    dict(runner.result("complete", ""), coverage={"scope": "text_layer", "total_pages": 1,
                         "visited_pages": 1, "pages_with_text": 0, "missing_text_pages": [[]]})]
        import json
        import subprocess
        for payload in payloads:
            process = unittest.mock.MagicMock()
            process.__enter__.return_value = process
            process.returncode = 0
            process.stdout.read.side_effect = [json.dumps(payload).encode(), b""]
            with patch.object(subprocess, "Popen", return_value=process):
                self.assertEqual(runner.extract(b"x", PDF)["reason"], "worker_protocol")

    def test_input_and_missing_interpreter(self):
        self.assertEqual(runner.extract(b"x", PDF, limit=0)["reason"], "invalid_input")
        self.assertEqual(runner.extract(b"xx", PDF, limit=1)["reason"], "size_limit")
        self.assertEqual(runner.extract(b"x", "unknown")["status"], "unsupported")
        self.assertEqual(runner.extract(b"x", PDF, optional_python="/missing/python")["status"], "unsupported")

    def test_clean_venv_has_no_dependencies(self):
        import subprocess
        with tempfile.TemporaryDirectory() as folder:
            subprocess.run([sys.executable, "-m", "venv", "--without-pip", folder], check=True)
            python = Path(folder) / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
            value = runner.extract(b"synthetic", PDF, optional_python=python)
            self.assertEqual(value["reason"], "missing_dependency")

    @unittest.skipIf(os.name == "nt", "synthetic executable fixture is POSIX")
    def test_process_and_output_bounds(self):
        # Executables simulate hostile workers; no parser mocking in acceptance tests.
        with tempfile.TemporaryDirectory() as folder:
            executable = Path(folder) / "worker"
            executable.write_text("#!/bin/sh\nexec /bin/sleep 5\n")
            executable.chmod(0o700)
            with patch.object(runner, "TIMEOUT", 0.05):
                self.assertEqual(runner.extract(b"x", PDF, optional_python=executable)["reason"], "process_timeout")
            executable.write_text("#!" + sys.executable + "\nimport sys\nsys.stdout.write('x'*100000)\n")
            with patch.object(runner, "MAX_OUTPUT", 1024):
                self.assertEqual(runner.extract(b"x", PDF, optional_python=executable)["reason"], "output_limit")
            executable.write_text("#!/bin/sh\nprintf '[]'\n")
            self.assertEqual(runner.extract(b"x", PDF, optional_python=executable)["reason"], "worker_protocol")


class ParserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            if any(worker.importlib.metadata.version(name) != version for name, version in worker.PINS.items()):
                raise unittest.SkipTest("Real parser suite requires exact requirements.lock pins")
            import pypdf, docx, openpyxl
        except (ImportError, worker.importlib.metadata.PackageNotFoundError):
            raise unittest.SkipTest("Real parser suite requires explicit requirements.lock provisioning")

    def run_extract(self, content, mime):
        original = hashlib.sha256(content).hexdigest()
        value = runner.extract(content, mime, optional_python=sys.executable)
        self.assertEqual(original, hashlib.sha256(content).hexdigest())
        self.assertEqual(value["output_hash"], hashlib.sha256(value["text"].encode()).hexdigest())
        self.assertTrue(runner.valid_coverage(value["coverage"]))
        for locator in value["locators"]:
            self.assertGreater(locator["end"], locator["start"])
            self.assertLessEqual(locator["end"], len(value["text"]))
        return value

    def pdf(self, text=True, encrypted=False):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject, NumberObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        if text:
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 12 Tf 10 100 Td (Synthetic page) Tj ET")
            page[NameObject("/Contents")] = writer._add_object(stream)
        else:
            # Image-only page, representing a synthetic scan with no text layer.
            image = DecodedStreamObject()
            image.set_data(b"\xff\xff\xff")
            for key, value in {"/Type": NameObject("/XObject"), "/Subtype": NameObject("/Image"),
                               "/Width": NumberObject(1), "/Height": NumberObject(1),
                               "/ColorSpace": NameObject("/DeviceRGB"), "/BitsPerComponent": NumberObject(8)}.items():
                image[NameObject(key)] = value
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Im1"): writer._add_object(image)})})
            stream = DecodedStreamObject()
            stream.set_data(b"q 100 0 0 100 0 0 cm /Im1 Do Q")
            page[NameObject("/Contents")] = writer._add_object(stream)
        if encrypted:
            writer.encrypt("synthetic-password")
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()

    def test_pdf_text_scanned_encrypted(self):
        value = self.run_extract(self.pdf(), PDF)
        self.assertEqual(value["status"], "complete")
        self.assertEqual(value["text"], "Synthetic page")
        self.assertEqual(value["locators"], [{"page": 1, "start": 0, "end": 14}])
        self.assertEqual(self.run_extract(self.pdf(False), PDF)["reason"], "no_text_layer")
        self.assertEqual(self.run_extract(self.pdf(encrypted=True), PDF)["status"], "encrypted")

    def test_mixed_pdf_coverage(self):
        from pypdf import PdfReader, PdfWriter
        writer = PdfWriter()
        for content in (self.pdf(), self.pdf(False)):
            writer.add_page(PdfReader(io.BytesIO(content)).pages[0])
        output = io.BytesIO()
        writer.write(output)
        value = self.run_extract(output.getvalue(), PDF)
        self.assertEqual(value["status"], "partial")
        self.assertEqual(value["coverage"], dict(scope="text_layer", total_pages=2,
                         visited_pages=2, pages_with_text=1, missing_text_pages=[2]))
        self.assertEqual(value["locators"], [{"page": 1, "start": 0, "end": 14}])

    def test_encoded_entities_and_ole_ambiguity(self):
        xml = '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "synthetic">]><x>&a;</x>'
        for encoding in ("utf-8-sig", "utf-16", "utf-16-be", "utf-32", "utf-32-be"):
            for mime in (DOCX, XLSX):
                value = self.run_extract(archive([("x.xml", xml.encode(encoding))]), mime)
                self.assertEqual(value["reason"], "unsafe_xml")
        value = self.run_extract(bytes.fromhex("d0cf11e0a1b11ae1") + b"synthetic", DOCX)
        self.assertEqual((value["status"], value["reason"]), ("failed", "encrypted_or_legacy_office"))

    def test_empty_documents(self):
        from pypdf import PdfWriter
        from docx import Document
        from openpyxl import Workbook
        for document, mime in ((PdfWriter(), PDF), (Document(), DOCX), (Workbook(), XLSX)):
            output = io.BytesIO()
            document.write(output) if mime == PDF else document.save(output)
            value = self.run_extract(output.getvalue(), mime)
            self.assertEqual(value["text"], "")
            self.assertEqual(value["locators"], [])
            self.assertEqual(value["status"], "partial" if mime == PDF else "complete")

    def test_renamed_xlsx_sheet_entity_rejected(self):
        from openpyxl import Workbook, load_workbook
        workbook = Workbook()
        workbook.active["A1"] = "Synthetic marker"
        source = io.BytesIO()
        workbook.save(source)
        # Both renamed paths are genuinely reachable via package relationships.
        for target in ("sheet1.XML", "sheet1.part"):
            entries = []
            with zipfile.ZipFile(io.BytesIO(source.getvalue())) as package:
                for name in package.namelist():
                    data = package.read(name)
                    if name == "xl/worksheets/sheet1.xml":
                        name = "xl/worksheets/" + target
                        data = (b'<!DOCTYPE worksheet [<!ENTITY synthetic "Expanded synthetic">]>'
                                + data.replace(b"Synthetic marker", b"&synthetic;"))
                    elif name in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                        data = data.replace(b"sheet1.xml", target.encode())
                    entries.append((name, data))
            content = archive(entries)
            # Prove the fixture reaches the actual parser and expands the entity
            # without preflight, rather than merely being a malformed ZIP.
            unguarded = load_workbook(io.BytesIO(content), read_only=True)
            try:
                self.assertEqual(unguarded.active["A1"].value, "Expanded synthetic")
            finally:
                unguarded.close()
            value = self.run_extract(content, XLSX)
            self.assertEqual((value["status"], value["reason"]), ("failed", "unsafe_xml"))
            self.assertEqual(value["text"], "")
            self.assertEqual(value["locators"], [])

    def test_docx_body_content_control_omission(self):
        from docx import Document
        from docx.oxml import OxmlElement
        document = Document()
        document.add_paragraph("Visible paragraph")
        control = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Synthetic controlled paragraph"
        run.append(text)
        paragraph.append(run)
        content.append(paragraph)
        control.append(content)
        document.element.body.insert(1, control)
        output = io.BytesIO()
        document.save(output)
        reopened = Document(io.BytesIO(output.getvalue()))
        self.assertEqual(len(reopened.element.xpath(".//w:sdt/w:sdtContent/w:p")), 1)
        self.assertEqual([p.text for p in reopened.paragraphs], ["Visible paragraph"])
        value = self.run_extract(output.getvalue(), DOCX)
        self.assertEqual(value["status"], "partial")
        self.assertEqual(value["coverage"]["omissions"]["content_controls"], 1)
        self.assertEqual(value["text"], "Visible paragraph")
        self.assertEqual(value["locators"], [{"paragraph": 1, "start": 0, "end": 17}])

    def test_docx_paragraphs_and_partial_table(self):
        from docx import Document
        document = Document()
        document.add_paragraph("First")
        document.add_paragraph("Second")
        output = io.BytesIO()
        document.save(output)
        value = self.run_extract(output.getvalue(), DOCX)
        self.assertEqual(value["status"], "complete")
        self.assertEqual(value["text"], "First\nSecond")
        self.assertEqual(value["coverage"]["scope"], "body_paragraphs")
        self.assertEqual(value["coverage"]["total_paragraphs"], 2)
        self.assertEqual(value["coverage"]["visited_paragraphs"], 2)
        self.assertEqual(value["locators"], [{"paragraph": 1, "start": 0, "end": 5}, {"paragraph": 2, "start": 6, "end": 12}])
        document.add_table(rows=1, cols=1)
        output = io.BytesIO()
        document.save(output)
        value = self.run_extract(output.getvalue(), DOCX)
        self.assertEqual(value["status"], "partial")
        self.assertEqual(value["coverage"]["omissions"]["tables"], 1)

    def test_xlsx_sheet_cells_formulas(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = "Synthetic"
        workbook.active["A1"] = "Alpha"
        workbook.active["C2"] = "=1+2"
        workbook.create_sheet("Other")["B1"] = 42
        output = io.BytesIO()
        workbook.save(output)
        value = self.run_extract(output.getvalue(), XLSX)
        self.assertEqual(value["status"], "complete")
        self.assertEqual(value["text"], "Alpha\n=1+2\n42")
        self.assertEqual(value["coverage"]["scope"], "stored_cells")
        self.assertEqual(value["coverage"]["visited_cells"], 3)
        self.assertEqual(value["coverage"]["nonempty_cells"], 3)
        self.assertEqual([(x["sheet"], x["cell"]) for x in value["locators"]], [("Synthetic", "A1"), ("Synthetic", "C2"), ("Other", "B1")])
        from openpyxl.comments import Comment
        workbook.active["A1"].comment = Comment("Synthetic note", "Synthetic author")
        output = io.BytesIO()
        workbook.save(output)
        value = self.run_extract(output.getvalue(), XLSX)
        self.assertEqual(value["status"], "partial")
        self.assertEqual(value["coverage"]["omissions"]["comments"], 1)

    def test_malformed_all_formats(self):
        for mime in (PDF, DOCX, XLSX):
            value = self.run_extract(b"malformed synthetic bytes", mime)
            self.assertEqual(value["reason"], "malformed_document")
            self.assertEqual(value["text"], "")

    def test_archive_guards(self):
        cases = [(archive([("../escape.xml", "x")]), "unsafe_archive_path"),
                 (archive([("/escape.xml", "x")]), "unsafe_archive_path"),
                 (archive([("x.xml", "x" * 1000000)], zipfile.ZIP_DEFLATED), "archive_limit"),
                 (archive([("x.xml", '<!DOCTYPE x [<!ENTITY a "x">]><x/>')]), "unsafe_xml")]
        encrypted = bytearray(archive([("x.xml", "x")]))
        for signature, offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
            position = encrypted.index(signature) + offset
            struct.pack_into("<H", encrypted, position, 1)
        cases.append((bytes(encrypted), "encrypted_archive"))
        for content, reason in cases:
            for mime in (DOCX, XLSX):
                self.assertEqual(self.run_extract(content, mime)["reason"], reason)


if __name__ == "__main__":
    unittest.main()
